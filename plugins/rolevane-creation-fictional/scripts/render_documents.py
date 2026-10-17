"""Format complete fictional host prose into new local drafts; never infer facts."""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import os
import re
import stat
import sys
from pathlib import Path
from urllib.parse import urlsplit

MAX_BYTES = 128 * 1024
LINK = re.compile(r"\[([^\[\]\n]+)\]\((https://[^\s()]+)\)")


def local_path(value: Path) -> Path:
    """Reject redirecting ancestors rather than silently resolving them."""
    path = Path(os.path.abspath(value))
    if path.drive.startswith("\\\\"):
        raise ValueError("UNC and device paths are unsupported; use host-local directories")
    for item in (path, *path.parents):
        try:
            info = item.lstat()
        except FileNotFoundError:
            continue
        if stat.S_ISLNK(info.st_mode) or getattr(info, "st_file_attributes", 0) & 0x400:
            raise ValueError("Symlink and reparse paths are unsupported")
    return path


def contained(value: Path, root: Path) -> Path:
    path, boundary = local_path(value), local_path(root)
    if not boundary.is_dir() or path == boundary or not path.is_relative_to(boundary):
        raise ValueError("Path must be inside an existing explicit root")
    return path


def inline(text: str) -> list[tuple[str, str | None]]:
    parts = []
    start = 0
    for match in LINK.finditer(text):
        parts.append((text[start : match.start()], None))
        url = match[2]
        if "\\" in url:
            raise ValueError("Hyperlinks cannot contain backslashes")
        parsed = urlsplit(url)
        host = parsed.hostname or ""
        if (
            not parsed.netloc.isascii()
            or len(host) > 253
            or any(
                not re.fullmatch(r"[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?", label)
                for label in host.split(".")
            )
            or parsed.netloc.endswith(":")
        ):
            raise ValueError("Hyperlinks require a valid ASCII DNS hostname and optional port")
        # Access validates the optional port's syntax and range before any files exist.
        _ = parsed.port
        if (
            parsed.username is not None
            or parsed.password is not None
            or not (
                host.endswith((".example", ".invalid", ".test"))
                or host in {"example.com", "example.org", "example.net"}
            )
        ):
            raise ValueError("Fictional hyperlinks require reserved example domains")
        parts.append((match[1], url))
        start = match.end()
    parts.append((text[start:], None))
    return parts


def read_blocks(path: Path) -> tuple[bytes, list[tuple[str, str]]]:
    if path.suffix.lower() != ".md" or not path.is_file():
        raise ValueError("Input must be one UTF-8 Markdown file")
    with path.open("rb") as source:
        raw = source.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        raise ValueError("Input exceeds 128 KiB")
    text = raw.decode("utf-8-sig")
    if any(ord(char) < 32 and char not in "\n\r\t" for char in text):
        raise ValueError("Control characters are unsupported")
    blocks: list[tuple[str, str]] = []
    pending: list[str] = []

    def flush() -> None:
        if pending:
            blocks.append(("paragraph", " ".join(pending)))
            pending.clear()

    for line in text.splitlines():
        line = line.strip()
        if len(line) > 2000:
            raise ValueError("Input line exceeds 2000 characters")
        if not line:
            flush()
        elif line.startswith(("# ", "## ", "- ")):
            flush()
            kind = "title" if line.startswith("# ") else "heading"
            if line.startswith("- "):
                kind = "bullet"
            blocks.append((kind, line.split(" ", 1)[1]))
        elif line.startswith(("#", "```", "~~~", "![", "<", "|", ">")):
            raise ValueError("Unsupported Markdown structure; use headings, paragraphs and bullets")
        else:
            pending.append(line)
    flush()
    if not blocks or blocks[0][0] != "title" or sum(k == "title" for k, _ in blocks) != 1:
        raise ValueError("Exactly one leading # title is required")
    if len(blocks) > 1000 or any(not text or len(text) > 8000 for _, text in blocks):
        raise ValueError("Empty or excessive content blocks")
    for _, value in blocks:
        inline(value)
    return raw, blocks


def pdf_layout(placements: list[dict], page_count: int) -> dict:
    """Diagnose this renderer's drawn blocks; this is not arbitrary PDF inspection."""
    pages, warnings = [], []
    for page_number in range(1, page_count + 1):
        items = [item for item in placements if item["page"] == page_number]
        codes = []
        if not items:
            pages.append({"page": page_number, "fragments": 0})
            warnings.append({"page": page_number, "code": "empty_page"})
            continue
        top = max(item["y"] + item["height"] for item in items)
        bottom = min(item["y"] for item in items)
        largest_gap = max(
            (a["y"] - b["y"] - b["height"] for a, b in zip(items, items[1:], strict=False)),
            default=0,
        )
        if any(
            item["x"] < 47.5
            or item["x"] + item["width"] > 564.5
            or item["y"] < 47.5
            or item["y"] + item["height"] > 744.5
            for item in items
        ):
            codes.append("outside_content_box")
        if items[-1]["kind"] in {"title", "heading"}:
            codes.append("orphan_heading")
        if largest_gap > 96:
            codes.append("large_internal_gap")
        if page_number < page_count and bottom - 48 > 144:
            codes.append("early_page_end")
        # A short single-page letter is valid; only a spill onto a sparse last page
        # warrants a warning. These thresholds belong to the fixed neutral layout.
        if page_count > 1 and page_number == page_count and top - bottom < 108:
            codes.append("thin_last_page")
        pages.append(
            {
                "page": page_number,
                "fragments": len(items),
                "top_pt": round(top, 2),
                "bottom_pt": round(bottom, 2),
                "largest_gap_pt": round(max(0, largest_gap), 2),
            }
        )
        warnings.extend({"page": page_number, "code": code} for code in codes)
    return {
        "status": "warnings" if warnings else "no_warnings",
        "basis": "renderer_drawn_blocks",
        "page_count": page_count,
        "content_box_pt": [48, 48, 564, 744],
        "pages": pages,
        "warnings": warnings,
    }


def render_pdf(path: Path, blocks: list[tuple[str, str]]) -> dict:
    import reportlab
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate
    from reportlab.platypus.flowables import KeepTogetherSplitAtTop

    placements = []

    class TrackedParagraph(Paragraph):
        kind = "paragraph"

        def split(self, availWidth, availHeight):
            fragments = super().split(availWidth, availHeight)
            for fragment in fragments:
                fragment.kind = self.kind
            return fragments

        def drawOn(self, canvas, x, y, _sW=0):
            super().drawOn(canvas, x, y, _sW)
            # wrap/split may run several times; only actual drawing is evidence.
            placements.append(
                {
                    "page": canvas.getPageNumber(),
                    "kind": self.kind,
                    "x": self._hAlignAdjust(x, _sW),
                    "y": y,
                    "width": self.width,
                    "height": self.height,
                }
            )

    fonts = Path(reportlab.__file__).parent / "fonts"
    for name, filename in (("DraftVera", "Vera.ttf"), ("DraftVeraBold", "VeraBd.ttf")):
        pdfmetrics.registerFont(TTFont(name, str(fonts / filename)))
    glyphs = pdfmetrics.getFont("DraftVera").face.charToGlyph
    if any(
        ord(char) not in glyphs for _, text in blocks for label, _ in inline(text) for char in label
    ):
        raise ValueError(
            "PDF font cannot represent a supplied character; use supported EN/DE prose"
        )
    styles = {
        "title": ParagraphStyle(
            "Title",
            fontName="DraftVeraBold",
            fontSize=19,
            leading=24,
            spaceAfter=16,
            keepWithNext=True,
        ),
        "heading": ParagraphStyle(
            "Heading",
            fontName="DraftVeraBold",
            fontSize=12,
            leading=16,
            spaceBefore=12,
            spaceAfter=7,
            keepWithNext=True,
        ),
        "paragraph": ParagraphStyle(
            "Body", fontName="DraftVera", fontSize=10, leading=15, spaceAfter=8, splitLongWords=True
        ),
        "bullet": ParagraphStyle(
            "Bullet",
            fontName="DraftVera",
            fontSize=10,
            leading=15,
            leftIndent=12,
            firstLineIndent=-10,
            spaceAfter=5,
            splitLongWords=True,
        ),
    }
    paragraphs = []
    for kind, text in blocks:
        content = "".join(
            f'<link href="{html.escape(url, quote=True)}" color="#245784">{html.escape(label)}</link>'
            if url
            else html.escape(label)
            for label, url in inline(text)
        )
        paragraph = TrackedParagraph(("• " if kind == "bullet" else "") + content, styles[kind])
        paragraph.kind = kind
        paragraphs.append(paragraph)
    story = []
    index = 0
    while index < len(blocks):
        end = index
        if blocks[index][0] == "paragraph":
            # Keep adjacent context/date paragraphs with their following bullet run.
            # Oversized blocks can split once they reach the top of a fresh page.
            while end < len(blocks) and blocks[end][0] == "paragraph":
                end += 1
            if end < len(blocks) and blocks[end][0] == "bullet":
                while end < len(blocks) and blocks[end][0] == "bullet":
                    end += 1
                story.append(KeepTogetherSplitAtTop(paragraphs[index:end]))
                index = end
                continue
        story.append(paragraphs[index])
        index += 1
    document = SimpleDocTemplate(
        str(path),
        # Uncompressed page streams are supported by the bounded Documents intake.
        # ReportLab's default ASCII85 compression is deliberately not accepted there.
        pageCompression=0,
        pagesize=(612, 792),
        leftMargin=48,
        rightMargin=48,
        topMargin=48,
        bottomMargin=48,
        title=blocks[0][1],
        author="",
        subject="Local unreviewed draft",
    )
    document.build(story)
    return pdf_layout(placements, document.page)


def render_docx(path: Path, blocks: list[tuple[str, str]]) -> None:
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(215.9), Mm(279.4)
    section.top_margin = section.bottom_margin = Mm(17)
    section.left_margin = section.right_margin = Mm(17)
    for name, size in (("Normal", 10), ("Title", 19), ("Heading 1", 12)):
        style = doc.styles[name]
        style.font.name, style.font.size = "Arial", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(8)
        if name == "Title":
            # The dependency's default Title style contains a blue bottom rule.
            border = style.element.get_or_add_pPr().find(qn("w:pBdr"))
            if border is not None:
                border.getparent().remove(border)
            style.font.underline = False
    doc.styles["Normal"].paragraph_format.line_spacing = 1.15
    doc.core_properties.title = blocks[0][1]
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.subject = "Local unreviewed draft"
    paragraphs = []
    for kind, text in blocks:
        style = {
            "title": "Title",
            "heading": "Heading 1",
            "paragraph": "Normal",
            "bullet": "List Bullet",
        }[kind]
        paragraph = doc.add_paragraph(style=style)
        paragraphs.append(paragraph)
        if kind in {"title", "heading"}:
            paragraph.paragraph_format.keep_with_next = True
        for label, url in inline(text):
            if not url:
                paragraph.add_run(label)
                continue
            link = OxmlElement("w:hyperlink")
            link.set(
                qn("r:id"),
                paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True),
            )
            run, props, color = OxmlElement("w:r"), OxmlElement("w:rPr"), OxmlElement("w:color")
            color.set(qn("w:val"), "245784")
            props.append(color)
            run.append(props)
            value = OxmlElement("w:t")
            value.text = label
            run.append(value)
            link.append(run)
            paragraph._p.append(link)
    index = 0
    while index < len(blocks):
        if blocks[index][0] != "paragraph":
            index += 1
            continue
        end = index
        while end < len(blocks) and blocks[end][0] == "paragraph":
            end += 1
        while end < len(blocks) and blocks[end][0] == "bullet":
            end += 1
        # Keep short context/date/qualification runs together. Long sections
        # retain normal pagination rather than creating an unbounded keep chain.
        if end - index <= 8 and sum(len(text) for _, text in blocks[index:end]) <= 1600:
            for paragraph in paragraphs[index:end]:
                paragraph.paragraph_format.keep_together = True
            for paragraph in paragraphs[index : end - 1]:
                paragraph.paragraph_format.keep_with_next = True
        index = end
    doc.save(path)


def create(
    input_root: Path, source: Path, output_root: Path, destination: Path, formats: list[str]
) -> dict:
    source = contained(source, input_root)
    destination = contained(destination, output_root)
    if (
        not formats
        or len(set(formats)) != len(formats)
        or any(f not in {"pdf", "docx"} for f in formats)
    ):
        raise ValueError("Select pdf and/or docx once each")
    if not destination.parent.is_dir():
        raise ValueError("Output parent must already exist")
    raw, blocks = read_blocks(source)
    # Capability checks happen before acquiring the exclusive destination.
    for name in formats:
        __import__("reportlab" if name == "pdf" else "docx")
    destination.mkdir(exist_ok=False)
    owned: list[Path] = []
    layout = None
    try:
        for name in formats:
            path = destination / f"draft.{name}"
            owned.append(path)
            if name == "pdf":
                layout = render_pdf(path, blocks)
            else:
                render_docx(path, blocks)
        return {
            "status": "local_unreviewed_draft",
            "source_sha256": hashlib.sha256(raw).hexdigest(),
            "files": [str(p) for p in owned],
            "visual_review": "required",
            **({"pdf_layout": layout} if layout is not None else {}),
        }
    except Exception:
        for path in owned:
            path.unlink(missing_ok=True)
        destination.rmdir()
        raise


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    for name in ("input-root", "input", "output-root", "output-dir"):
        parser.add_argument(f"--{name}", type=Path, required=True)
    parser.add_argument("--format", choices=("pdf", "docx"), action="append", required=True)
    args = parser.parse_args()
    try:
        receipt = create(
            args.input_root, args.input, args.output_root, args.output_dir, args.format
        )
    except (OSError, ValueError, ImportError) as error:
        print(f"Rendering failed: {error}", file=sys.stderr)
        return 2
    print(json.dumps(receipt, ensure_ascii=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
