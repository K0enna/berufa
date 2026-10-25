"""Protect layout feedback at the real renderer boundary, with actual PDF readback."""

import importlib.util
import re
import textwrap
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def renderer():
    pytest.importorskip("reportlab")
    spec = importlib.util.spec_from_file_location(
        "creation_renderer",
        ROOT / "plugins/rolevane-creation-fictional/scripts/render_documents.py",
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.mark.parametrize(
    ("body", "expected_warning"),
    [
        ("A short fictional application answer.", None),
        ("A grounded experience statement.\n\n## Section without content", "orphan_heading"),
        (
            "\n\n".join(
                f"Entry {i}: Coached colleagues and reviewed service cases." for i in range(29)
            ),
            "thin_last_page",
        ),
    ],
)
def test_rendered_pdf_reports_actionable_layout_without_rewriting(
    tmp_path, renderer, body, expected_warning
):
    from pypdf import PdfReader

    source = tmp_path / "source.md"
    text = f"# Fictional CV\n\n{body}\n"
    source.write_text(text, encoding="utf-8")
    result = renderer.create(tmp_path, source, tmp_path, tmp_path / "draft", ["pdf"])
    pdf = PdfReader(result["files"][0])
    layout = result["pdf_layout"]
    assert layout["page_count"] == len(pdf.pages)
    assert len(layout["pages"]) == len(pdf.pages)
    warnings = {warning["code"] for warning in layout["warnings"]}
    assert warnings == ({expected_warning} if expected_warning else set())
    assert result["status"] == "local_unreviewed_draft"
    assert result["visual_review"] == "required"
    extracted = " ".join(" ".join(page.extract_text().split()) for page in pdf.pages)
    for line in body.splitlines():
        assert line.removeprefix("## ") in extracted
    assert source.read_text(encoding="utf-8") == text
    original = Path(result["files"][0]).read_bytes()
    with pytest.raises(FileExistsError):
        renderer.create(tmp_path, source, tmp_path, tmp_path / "draft", ["pdf"])
    assert Path(result["files"][0]).read_bytes() == original


def test_long_paragraph_records_actual_split_pages(tmp_path, renderer):
    from pypdf import PdfReader

    source = tmp_path / "source.md"
    # One permitted block spans pages. Counting wrap probes or losing split
    # fragments would disagree with actual rendered page/text evidence.
    prose = " ".join(f"case{i} discussed with colleagues" for i in range(210))
    source.write_text(f"# Fictional CV\n\n{textwrap.fill(prose, width=1000)}", encoding="utf-8")
    result = renderer.create(tmp_path, source, tmp_path, tmp_path / "draft", ["pdf"])
    pdf = PdfReader(result["files"][0])
    assert len(pdf.pages) > 1
    layout = result["pdf_layout"]
    assert layout["page_count"] == len(pdf.pages)
    assert all(page["fragments"] > 0 for page in layout["pages"])
    assert not {"outside_content_box", "empty_page", "orphan_heading"} & {
        warning["code"] for warning in layout["warnings"]
    }
    extracted = " ".join(page.extract_text() for page in pdf.pages)
    assert re.findall(r"\bcase(\d+)\b", extracted) == [str(index) for index in range(210)]


@pytest.mark.parametrize("title", ["Fictional letter", "A" * 255, "A" * 256])
def test_docx_only_receipt_has_no_pdf_verification_claim(tmp_path, renderer, title):
    pytest.importorskip("docx")
    from docx import Document

    source = tmp_path / "source.md"
    source.write_text(f"# {title}\n\nCoached colleagues.", encoding="utf-8")
    result = renderer.create(tmp_path, source, tmp_path, tmp_path / "draft", ["docx"])
    assert "pdf_layout" not in result
    assert result["visual_review"] == "required"
    document = Document(result["files"][0])
    assert "Coached colleagues." in [p.text for p in document.paragraphs]
    assert document.paragraphs[0].text == title
    assert document.core_properties.title == title[:255]
